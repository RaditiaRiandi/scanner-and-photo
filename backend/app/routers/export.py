import io
import datetime
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..database import get_db
from ..models import HartaScan, AntamPhoto

router = APIRouter(prefix="/api/export", tags=["Export Excel"])

def make_header_style():
    fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    return fill, font, border

def make_row_border():
    return Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

def autofit(ws, headers, data_row_count):
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1]))
        for row in range(2, data_row_count + 2):
            val = ws.cell(row=row, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

def write_harta_sheet(ws, items):
    headers = ["No", "Waktu Scan", "Nama Pengguna (Device)", "Informasi Harta (Link URL)", "ID Mandiri"]
    ws.append(headers)
    hfill, hfont, hborder = make_header_style()
    rborder = make_row_border()
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = hfill; c.font = hfont
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = hborder
    ws.row_dimensions[1].height = 28

    for idx, item in enumerate(items, 1):
        loc_t = item.scanned_at + datetime.timedelta(hours=7)
        rn = idx + 1
        ws.cell(rn, 1, idx).alignment = Alignment(horizontal="center")
        ws.cell(rn, 2, loc_t.strftime("%Y-%m-%d %H:%M:%S")).alignment = Alignment(horizontal="center")
        ws.cell(rn, 3, item.device_id).alignment = Alignment(horizontal="left")
        c_url = ws.cell(rn, 4, item.url_harta)
        c_url.alignment = Alignment(horizontal="left")
        if item.url_harta.startswith("http"):
            c_url.font = Font(name="Calibri", size=10, color="0284C7", underline="single")
            c_url.hyperlink = item.url_harta
        c_id = ws.cell(rn, 5, item.id_mandiri)
        c_id.font = Font(name="Consolas", size=10)
        for ci in range(1, 6):
            ws.cell(rn, ci).border = rborder
        ws.row_dimensions[rn].height = 22
    autofit(ws, headers, len(items))

def write_antam_sheet(ws, items):
    headers = ["No", "Waktu Upload", "Nama Pengguna (Device)", "Gramasi (g)", "Nama File", "Folder Penyimpanan"]
    ws.append(headers)
    hfill, hfont, hborder = make_header_style()
    rborder = make_row_border()
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = hfill; c.font = hfont
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = hborder
    ws.row_dimensions[1].height = 28

    for idx, item in enumerate(items, 1):
        loc_t = item.uploaded_at + datetime.timedelta(hours=7)
        rn = idx + 1
        ws.cell(rn, 1, idx).alignment = Alignment(horizontal="center")
        ws.cell(rn, 2, loc_t.strftime("%Y-%m-%d %H:%M:%S")).alignment = Alignment(horizontal="center")
        ws.cell(rn, 3, item.device_id).alignment = Alignment(horizontal="left")
        g_val = int(item.gramasi) if item.gramasi and item.gramasi.isdigit() else 0
        c_g = ws.cell(rn, 4, g_val)
        c_g.alignment = Alignment(horizontal="center")
        c_g.font = Font(name="Calibri", bold=True, size=10)
        ws.cell(rn, 5, item.filename).alignment = Alignment(horizontal="left")
        ws.cell(rn, 6, item.filepath).alignment = Alignment(horizontal="left")
        for ci in range(1, 7):
            ws.cell(rn, ci).border = rborder
        ws.row_dimensions[rn].height = 22
    autofit(ws, headers, len(items))

def write_gramasi_summary_sheet(ws, items):
    """Sheet ringkasan per gramasi."""
    from collections import defaultdict
    summary = defaultdict(lambda: {"count": 0, "total_gram": 0})
    for item in items:
        g = item.gramasi or "unknown"
        g_val = int(item.gramasi) if item.gramasi and item.gramasi.isdigit() else 0
        summary[g]["count"] += 1
        summary[g]["total_gram"] += g_val

    headers = ["Gramasi (g)", "Jumlah Foto", "Total Berat (g)", "Proporsi (%)"]
    ws.append(headers)
    hfill, hfont, hborder = make_header_style()
    rborder = make_row_border()
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = hfill; c.font = hfont
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = hborder
    ws.row_dimensions[1].height = 28

    total_count = sum(v["count"] for v in summary.values())
    sorted_keys = sorted(summary.keys(), key=lambda x: int(x) if x.isdigit() else 9999)
    for idx, g in enumerate(sorted_keys, 1):
        rn = idx + 1
        d = summary[g]
        pct = round(d["count"] / total_count * 100, 1) if total_count else 0
        ws.cell(rn, 1, f"{g}g").alignment = Alignment(horizontal="center")
        ws.cell(rn, 2, d["count"]).alignment = Alignment(horizontal="center")
        ws.cell(rn, 3, d["total_gram"]).alignment = Alignment(horizontal="center")
        ws.cell(rn, 4, f"{pct}%").alignment = Alignment(horizontal="center")
        for ci in range(1, 5):
            ws.cell(rn, ci).border = rborder
        ws.row_dimensions[rn].height = 22

    # Total row
    total_rn = len(sorted_keys) + 2
    fill_total = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    font_total = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    ws.cell(total_rn, 1, "TOTAL").font = font_total
    ws.cell(total_rn, 1).fill = fill_total
    ws.cell(total_rn, 1).alignment = Alignment(horizontal="center")
    ws.cell(total_rn, 2, total_count).font = font_total
    ws.cell(total_rn, 2).fill = fill_total
    ws.cell(total_rn, 2).alignment = Alignment(horizontal="center")
    total_g = sum(v["total_gram"] for v in summary.values())
    ws.cell(total_rn, 3, total_g).font = font_total
    ws.cell(total_rn, 3).fill = fill_total
    ws.cell(total_rn, 3).alignment = Alignment(horizontal="center")
    ws.cell(total_rn, 4, "100%").font = font_total
    ws.cell(total_rn, 4).fill = fill_total
    ws.cell(total_rn, 4).alignment = Alignment(horizontal="center")
    ws.row_dimensions[total_rn].height = 24

    autofit(ws, headers, len(sorted_keys) + 1)

@router.get("/harta")
def export_harta(db: Session = Depends(get_db)):
    items = db.query(HartaScan).order_by(HartaScan.id.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Scan Harta"
    write_harta_sheet(ws, items)
    output = io.BytesIO()
    wb.save(output); output.seek(0)
    fname = f"Data_Scan_Harta_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )

@router.get("/antam")
def export_antam(db: Session = Depends(get_db)):
    items = db.query(AntamPhoto).order_by(AntamPhoto.id.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Foto Antam"
    write_antam_sheet(ws, items)
    output = io.BytesIO()
    wb.save(output); output.seek(0)
    fname = f"Data_Foto_Antam_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )

@router.get("/all")
def export_all(db: Session = Depends(get_db)):
    h_items = db.query(HartaScan).order_by(HartaScan.id.asc()).all()
    a_items = db.query(AntamPhoto).order_by(AntamPhoto.id.asc()).all()

    wb = openpyxl.Workbook()

    # Sheet 1: Harta
    ws1 = wb.active
    ws1.title = "Rekap Harta"
    write_harta_sheet(ws1, h_items)

    # Sheet 2: Antam Detail
    ws2 = wb.create_sheet(title="Rekap Antam")
    write_antam_sheet(ws2, a_items)

    # Sheet 3: Ringkasan Gramasi
    ws3 = wb.create_sheet(title="Ringkasan Gramasi")
    write_gramasi_summary_sheet(ws3, a_items)

    output = io.BytesIO()
    wb.save(output); output.seek(0)
    fname = f"Rekap_Lengkap_PlanC_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )
